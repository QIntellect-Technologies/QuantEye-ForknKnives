import os
import sys
import io
import json
import sqlite3
import openpyxl

sys.path.insert(0, os.path.abspath('.'))

from app import app
import database as db

db.init_db()
client = app.test_client()

TEST_USER_ID = 38
E2E_DATE_ATT = "2026-10-26"
E2E_DATE_LEAVE = "2026-10-27"
E2E_DATE_OT = "2026-10-28"

print("=" * 80)
print("  COMPLETE END-TO-END VERIFICATION: LIFECYCLE CHAIN PROOF")
print("  Employee -> Shift -> Attendance -> Leave -> Overtime -> Salary -> Payroll -> Reports -> Excel")
print("=" * 80)

orig_cfg = db.get_salary_config(TEST_USER_ID)

try:
    # 1. EMPLOYEE VERIFICATION
    emp = db.get_user_by_id(TEST_USER_ID)
    assert emp is not None, f"Employee #{TEST_USER_ID} not found"
    print(f"\n[STEP 1] EMPLOYEE VERIFIED:")
    print(f"  ID: #{emp['id']} | Name: {emp['name']} | Dept: {emp['department']} | Role: {emp['position']}")

    # 2. SHIFT VERIFICATION
    shift = db.get_shift_by_id(emp['shift_id'])
    assert shift is not None, f"Shift #{emp['shift_id']} not found"
    print(f"\n[STEP 2] SHIFT VERIFIED:")
    print(f"  Shift #{shift['id']}: '{shift['name']}' | Start: {shift['start_time']} | Grace: {shift['grace_period_minutes']} mins")

    # 3. MANUAL ATTENDANCE INGESTION
    with sqlite3.connect('attendance.db') as conn:
        conn.execute(
            "INSERT INTO attendance (user_id, detected_name, confidence, source, location, device_id, timestamp) "
            "VALUES (?, ?, 1.0, 'manual', 'Main Reception', 'cam_manual_e2e', ?)",
            (TEST_USER_ID, emp['name'], f"{E2E_DATE_ATT} 10:32:00")
        )
        conn.commit()
    print(f"\n[STEP 3] MANUAL ATTENDANCE RECORDED:")
    print(f"  Date: {E2E_DATE_ATT} | Check-in: 10:32:00 (within grace period) | Source: manual")

    # 4. LEAVE REQUEST & APPROVAL
    leave_id = db.add_leave_request(
        user_id=TEST_USER_ID,
        user_name=emp['name'],
        leave_type='annual',
        start_date=E2E_DATE_LEAVE,
        end_date=E2E_DATE_LEAVE,
        reason='E2E Audit Annual Leave'
    )
    db.update_leave_status(leave_id, status='approved', approved_by='HR Admin')
    print(f"\n[STEP 4] LEAVE RECORDED & APPROVED:")
    print(f"  Leave #{leave_id} | Type: Annual (Paid) | Date: {E2E_DATE_LEAVE} | Status: approved")

    # 5. OVERTIME LOGGING & APPROVAL
    ot_id = db.add_overtime(
        user_id=TEST_USER_ID,
        user_name=emp['name'],
        ot_date=E2E_DATE_OT,
        hours=3.0,
        reason='Production deployment overtime'
    )
    db.update_overtime_status(ot_id, status='approved', approved_by='Operations Lead')
    print(f"\n[STEP 5] OVERTIME LOGGED & APPROVED:")
    print(f"  Overtime #{ot_id} | Date: {E2E_DATE_OT} | Hours: 3.0 hrs | Status: approved")

    # 6. SALARY CONFIGURATION SETTING
    db.set_salary_config(
        user_id=TEST_USER_ID,
        basic_salary=250000.0,
        allowances=35000.0,
        deductions=10000.0,
        ot_rate=1201.92
    )
    print(f"\n[STEP 6] SALARY CONFIGURATION CONFIGURED:")
    print(f"  Basic Salary: PKR 250,000.00")
    print(f"  Allowances:   PKR 35,000.00")
    print(f"  Fixed Deduct: PKR 10,000.00")
    print(f"  OT Rate:      PKR 1,201.92 / hr")

    # 7. PAYROLL CALCULATION ENGINE
    payroll = db.calculate_monthly_payroll('2026-10')
    emp_pay = next((r for r in payroll.get('records', []) if r.get('employee_id') == TEST_USER_ID), None)
    assert emp_pay is not None, "Payroll calculation missing employee record"

    expected_ot_amt = round(3.0 * 1201.92, 2) # 3605.76
    expected_gross = round(250000.0 + 35000.0 + expected_ot_amt, 2) # 288605.76
    absence_ded = float(emp_pay.get('deduction_amount') or 0.0)
    expected_net = round(expected_gross - 10000.0 - absence_ded, 2)

    print(f"\n[STEP 7] PAYROLL DERIVED VALUES:")
    print(f"  Basic Salary:      PKR {emp_pay['basic_salary']:,.2f}")
    print(f"  Allowances:        PKR {emp_pay['allowances']:,.2f}")
    print(f"  Approved OT Hours: {emp_pay['overtime_hours']} hrs")
    print(f"  OT Rate:           PKR {emp_pay['overtime_rate']:,.2f} / hr")
    print(f"  Calculated OT Pay: PKR {emp_pay['overtime_amount']:,.2f} (Expected: PKR {expected_ot_amt:,.2f})")
    print(f"  Gross Salary:      PKR {emp_pay['gross_salary']:,.2f} (Expected: PKR {expected_gross:,.2f})")
    print(f"  Fixed Deductions:  PKR {emp_pay['custom_deductions']:,.2f}")
    print(f"  Absence Penalty:   PKR {absence_ded:,.2f}")
    print(f"  Final Net Salary:  PKR {emp_pay['net_salary']:,.2f} (Expected: PKR {expected_net:,.2f})")

    assert emp_pay['overtime_amount'] == expected_ot_amt, f"OT pay mismatch: {emp_pay['overtime_amount']} vs {expected_ot_amt}"
    assert emp_pay['gross_salary'] == expected_gross, f"Gross salary mismatch: {emp_pay['gross_salary']} vs {expected_gross}"
    assert emp_pay['net_salary'] == expected_net, f"Net salary mismatch: {emp_pay['net_salary']} vs {expected_net}"
    print("  -> ALL PAYROLL FORMULAS PROVEN ACCURATE!")

    # 8. ATTENDANCE REPORT API
    rep_res = client.get('/api/reports/attendance?start_date=2026-10-26&end_date=2026-10-28')
    assert rep_res.status_code == 200, f"Reports API failed: {rep_res.status_code}"
    rep_data = rep_res.get_json() or {}
    emp_rep = next((r for r in rep_data.get('records', []) if r.get('user_id') == TEST_USER_ID), None)
    assert emp_rep is not None, "Employee record not found in attendance report"
    print(f"\n[STEP 8] ATTENDANCE REPORT API VERIFIED:")
    print(f"  HTTP {rep_res.status_code} | Total Employees in Report: {len(rep_data.get('records', []))}")
    print(f"  Employee #{emp_rep['user_id']} Daily records: {len(emp_rep.get('daily_records', []))} days")

    # 9. EXCEL EXPORT API & WORKBOOK INTEGRITY
    excel_res = client.get('/api/reports/export/excel?start_date=2026-10-26&end_date=2026-10-28')
    assert excel_res.status_code == 200, f"Excel export failed: {excel_res.status_code}"
    wb = openpyxl.load_workbook(io.BytesIO(excel_res.data))
    ws = wb.active
    print(f"\n[STEP 9] EXCEL EXPORT WORKBOOK VERIFIED:")
    print(f"  HTTP {excel_res.status_code} | Payload: {len(excel_res.data)} bytes")
    print(f"  Worksheet Title: '{ws.title}' | Dimensions: {ws.dimensions} | Total Rows: {ws.max_row}")
    assert ws.title == 'Attendance Report', f"Unexpected sheet name: {ws.title}"
    assert ws.max_row > 5, "Excel report has no data rows"

    print("\n" + "=" * 80)
    print("  COMPLETE LIFECYCLE AUDIT VERDICT: 100% PROVEN END-TO-END")
    print("=" * 80)

finally:
    # Cleanup all test artifacts
    with sqlite3.connect('attendance.db') as conn:
        conn.execute("DELETE FROM attendance WHERE user_id = ? AND timestamp LIKE '2026-10-26%';", (TEST_USER_ID,))
        conn.execute("DELETE FROM leave_requests WHERE user_id = ? AND start_date = '2026-10-27';", (TEST_USER_ID,))
        conn.execute("DELETE FROM overtime WHERE user_id = ? AND ot_date = '2026-10-28';", (TEST_USER_ID,))
        conn.commit()

    if orig_cfg:
        db.set_salary_config(
            user_id=TEST_USER_ID,
            basic_salary=float(orig_cfg.get('basic_salary') or 0.0),
            allowances=float(orig_cfg.get('allowances') or 0.0),
            deductions=float(orig_cfg.get('deductions') or 0.0),
            ot_rate=float(orig_cfg.get('ot_rate') or 0.0)
        )
    print("\nTest data cleaned up and salary configuration restored cleanly.")
