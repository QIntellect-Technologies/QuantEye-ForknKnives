import io
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_attendance_excel(report_data, filter_info=None):
    """
    Generate a beautifully styled Excel report matching the active report data.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"
    ws.views.sheetView[0].showGridLines = True

    # Color tokens
    primary_color = "1E3A8A"       # Deep Navy Blue #1E3A8A
    header_fill = PatternFill(start_color=primary_color, end_color=primary_color, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    title_font = Font(name="Calibri", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    meta_label_font = Font(name="Calibri", size=9, bold=True, color="475569")
    meta_val_font = Font(name="Calibri", size=9, color="0F172A")
    
    stat_box_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border_side = Side(style='thin', color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # Status Pill Colors
    present_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    present_font = Font(name="Calibri", size=10, bold=True, color="065F46")
    
    late_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    late_font = Font(name="Calibri", size=10, bold=True, color="92400E")
    
    absent_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    absent_font = Font(name="Calibri", size=10, bold=True, color="991B1B")

    # 1. Document Title
    ws.merge_cells('A1:J1')
    title_cell = ws['A1']
    title_cell.value = "QIntellect Technologies - Official Attendance Report"
    title_cell.font = title_font
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    # 2. Metadata subtitle & active filters
    filter_info = filter_info or {}
    date_label = report_data.get('date_label', 'Today')
    dept_label = filter_info.get('department') or 'All Departments'
    shift_label = filter_info.get('shift') or 'All Shifts'
    status_label = filter_info.get('status') or 'All Statuses'
    
    ws.merge_cells('A2:J2')
    sub_cell = ws['A2']
    sub_cell.value = f"Report Period: {date_label}  |  Department: {dept_label}  |  Shift: {shift_label}  |  Status Filter: {status_label}"
    sub_cell.font = subtitle_font
    sub_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 18

    # 3. Summary KPI Box
    summary = report_data.get('summary', {})
    tot = summary.get('total_registered', 0)
    pres = summary.get('present', 0)
    absn = summary.get('absent', 0)
    late = summary.get('late', 0)
    rate = summary.get('attendance_rate', 0.0)

    ws['A4'] = "Total Records"
    ws['A5'] = tot
    ws['C4'] = "Present"
    ws['C5'] = pres
    ws['E4'] = "Late"
    ws['E5'] = late
    ws['G4'] = "Absent"
    ws['G5'] = absn
    ws['I4'] = "Attendance Rate"
    ws['I5'] = f"{rate}%"

    for col in ['A', 'C', 'E', 'G', 'I']:
        cell_lbl = ws[f"{col}4"]
        cell_lbl.font = meta_label_font
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.fill = stat_box_fill
        cell_lbl.border = thin_border
        
        cell_val = ws[f"{col}5"]
        cell_val.font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.fill = stat_box_fill
        cell_val.border = thin_border

    ws.row_dimensions[4].height = 18
    ws.row_dimensions[5].height = 22

    # 4. Table Header Row (Row 7)
    headers = [
        ("Employee ID", 14),
        ("Employee Name", 24),
        ("Department", 18),
        ("Position / Role", 18),
        ("Shift", 20),
        ("Date", 14),
        ("Status", 14),
        ("Check-In Time", 16),
        ("Check-Out Time", 16),
        ("Location / Camera", 20),
        ("AI Confidence", 14)
    ]

    header_row = 7
    ws.row_dimensions[header_row].height = 26

    for col_idx, (header_text, width) in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    # 5. Data Rows (Row 8+)
    records = report_data.get('records', [])
    for row_idx, r in enumerate(records, start=header_row + 1):
        ws.row_dimensions[row_idx].height = 20
        is_alt = (row_idx % 2 == 0)
        current_fill = alt_row_fill if is_alt else PatternFill(fill_type=None)

        st = (r.get('status') or 'absent').lower()
        if st in ('present', 'checked_out'):
            st_text = "PRESENT" if st == 'present' else "CHECKED OUT"
            st_font = present_font
            st_fill = present_fill
        elif st == 'late':
            st_text = "LATE"
            st_font = late_font
            st_fill = late_fill
        else:
            st_text = "ABSENT"
            st_font = absent_font
            st_fill = absent_fill

        first_seen = r.get('first_seen') or '--'
        last_seen = r.get('last_seen') or '--'
        conf_str = f"{r['confidence'] * 100:.1f}%" if r.get('confidence') else '--'
        shift_str = f"{r.get('shift_name', 'General')} ({r.get('shift_start', '08:00')}-{r.get('shift_end', '20:00')})"

        row_data = [
            (f"#{r.get('user_id', '')}", Alignment(horizontal="center"), Font(name="Calibri", size=10, bold=True, color="247EAE")),
            (r.get('name', ''), Alignment(horizontal="left"), Font(name="Calibri", size=10, bold=True, color="0F172A")),
            (r.get('department', 'Unassigned'), Alignment(horizontal="left"), Font(name="Calibri", size=10, color="334155")),
            (r.get('position', 'Staff Member'), Alignment(horizontal="left"), Font(name="Calibri", size=10, color="64748B")),
            (shift_str, Alignment(horizontal="left"), Font(name="Calibri", size=9, color="475569")),
            (r.get('date', ''), Alignment(horizontal="center"), Font(name="Calibri", size=10, color="0F172A")),
            (st_text, Alignment(horizontal="center"), st_font, st_fill),
            (first_seen, Alignment(horizontal="center"), Font(name="Calibri", size=10, color="0F172A")),
            (last_seen, Alignment(horizontal="center"), Font(name="Calibri", size=10, color="64748B")),
            (r.get('location') or 'Main Camera', Alignment(horizontal="left"), Font(name="Calibri", size=9, color="475569")),
            (conf_str, Alignment(horizontal="center"), Font(name="Calibri", size=9, bold=True, color="059669"))
        ]

        for col_idx, item in enumerate(row_data, 1):
            val = item[0]
            align = item[1]
            font = item[2]
            fill = item[3] if len(item) > 3 else current_fill

            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = align
            cell.font = font
            cell.border = thin_border
            if fill and fill.fill_type:
                cell.fill = fill

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

if __name__ == "__main__":
    import scratch.test_report_fix as tr
    rep = tr.test_new_report(target_date='yesterday')
    buf = generate_attendance_excel(rep, {'department': 'All', 'shift': 'All', 'status': 'All'})
    print("Excel size bytes:", len(buf.getvalue()))
